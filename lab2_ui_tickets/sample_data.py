from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook


APP_DIR = Path(__file__).resolve().parent


def create_students(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ИС-21"
    sheet.append(["Фамилия", "Имя"])
    sheet.append(["Иванов", "Иван"])
    sheet.append(["Петров", "Петр"])
    sheet.append(["Сидорова", "Анна"])

    sheet = workbook.create_sheet("ИС-22")
    sheet.append(["Фамилия", "Имя"])
    sheet.append(["Кузнецова", "Мария"])
    sheet.append(["Смирнов", "Алексей"])

    empty_sheet = workbook.create_sheet("Пустая группа")
    empty_sheet.append(["Фамилия", "Имя"])

    workbook.save(path)


def create_tickets(path: Path) -> None:
    document = Document()
    for ticket_number in range(1, 6):
        document.add_paragraph(f"Билет {ticket_number}")
        document.add_paragraph(f"1. Теоретический вопрос билета {ticket_number}")
        document.add_paragraph(f"2. Практическое задание билета {ticket_number}")
        document.add_paragraph(f"3. Дополнительный вопрос билета {ticket_number}")
        document.add_paragraph("")
    document.save(path)


if __name__ == "__main__":
    create_students(APP_DIR / "students.xlsx")
    create_tickets(APP_DIR / "tickets.docx")
    print("Созданы students.xlsx и tickets.docx")
