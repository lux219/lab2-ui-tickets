import logging

from docx import Document

from ticket_parser import parse_tickets


def test_parse_tickets_skips_invalid_tickets(tmp_path, caplog):
    docx_path = tmp_path / "tickets.docx"
    document = Document()
    document.add_paragraph("Билет 1")
    document.add_paragraph("1. Первый вопрос")
    document.add_paragraph("2. Второй вопрос")
    document.add_paragraph("3. Третий вопрос")
    document.add_paragraph("Билет 2")
    document.add_paragraph("1. Только один вопрос")
    document.add_paragraph("Билет 3")
    document.add_paragraph("1. Другой первый вопрос")
    document.add_paragraph("2. Другой второй вопрос")
    document.add_paragraph("3. Другой третий вопрос")
    document.save(docx_path)

    with caplog.at_level(logging.WARNING):
        tickets = parse_tickets(docx_path)

    assert sorted(tickets) == [1, 3]
    assert tickets[1].questions == ("Первый вопрос", "Второй вопрос", "Третий вопрос")
    assert "Билет 2 пропущен" in caplog.text


def test_parse_tickets_logs_unrecognized_header(tmp_path, caplog):
    docx_path = tmp_path / "tickets.docx"
    document = Document()
    document.add_paragraph("Билет A")
    document.add_paragraph("1. Вопрос без корректного номера билета")
    document.save(docx_path)

    with caplog.at_level(logging.WARNING):
        tickets = parse_tickets(docx_path)

    assert tickets == {}
    assert "Заголовок билета не распознан" in caplog.text
