import random
from datetime import datetime

from openpyxl import load_workbook

from results_store import HEADERS, ResultsStore
from student_store import Student


def test_repeated_student_gets_first_ticket_and_new_history_row(tmp_path):
    results_path = tmp_path / "results.xlsx"
    store = ResultsStore(results_path)
    student = Student(last_name="Иванов", first_name="Иван")

    first = store.record_assignment(
        "ИС-21",
        student,
        [1, 2, 3],
        now=lambda: datetime(2026, 9, 15, 10, 0, 0),
        rng=random.Random(1),
    )
    second = store.record_assignment(
        "ИС-21",
        student,
        [1, 2, 3],
        now=lambda: datetime(2026, 9, 15, 10, 5, 0),
        rng=random.Random(999),
    )

    assert first.ticket_number == second.ticket_number
    assert first.is_repeat is False
    assert second.is_repeat is True

    workbook = load_workbook(results_path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    assert rows[0] == HEADERS
    assert len(rows) == 3
    assert rows[1][3] == rows[2][3]
    assert rows[1][5] == "нет"
    assert rows[2][5] == "да"
