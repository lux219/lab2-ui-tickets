from __future__ import annotations

import random
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from collections.abc import Sequence
from typing import Callable

from openpyxl import Workbook, load_workbook

from student_store import Student


HEADERS = ("Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор (да/нет)")


class ResultsWriteError(RuntimeError):
    pass


@dataclass(frozen=True)
class Assignment:
    ticket_number: int
    is_repeat: bool


class ResultsStore:
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def record_assignment(
        self,
        group_name: str,
        student: Student,
        valid_ticket_numbers: Sequence[int],
        now: Callable[[], datetime] = datetime.now,
        rng: random.Random | None = None,
    ) -> Assignment:
        if not valid_ticket_numbers:
            raise ValueError("Нет доступных билетов для генерации.")

        random_source = rng or random
        workbook = self._open_or_create_workbook()
        sheet = workbook.active

        existing_ticket = self._find_first_ticket(sheet, group_name, student)
        is_repeat = existing_ticket is not None
        ticket_number = existing_ticket if existing_ticket is not None else random_source.choice(list(valid_ticket_numbers))

        sheet.append(
            [
                group_name,
                student.last_name,
                student.first_name,
                ticket_number,
                now().strftime("%Y-%m-%d %H:%M:%S"),
                "да" if is_repeat else "нет",
            ]
        )

        try:
            workbook.save(self.path)
        except PermissionError as error:
            raise ResultsWriteError(
                f"Не удалось записать журнал результатов. Возможно, файл открыт в Excel: {self.path}"
            ) from error
        except OSError as error:
            raise ResultsWriteError(f"Не удалось сохранить журнал результатов: {self.path}") from error
        finally:
            workbook.close()

        return Assignment(ticket_number=ticket_number, is_repeat=is_repeat)

    def _open_or_create_workbook(self) -> Workbook:
        if self.path.exists():
            workbook = load_workbook(self.path)
            sheet = workbook.active
            if sheet.max_row == 0 or [sheet.cell(1, index).value for index in range(1, 7)] != list(HEADERS):
                if sheet.max_row == 1 and all(sheet.cell(1, index).value is None for index in range(1, 7)):
                    for index, header in enumerate(HEADERS, start=1):
                        sheet.cell(1, index).value = header
            return workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"
        sheet.append(list(HEADERS))
        return workbook

    @staticmethod
    def _find_first_ticket(sheet, group_name: str, student: Student) -> int | None:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            group, last_name, first_name, ticket_number = row[:4]
            if (
                str(group).strip() == group_name
                and str(last_name).strip() == student.last_name
                and str(first_name).strip() == student.first_name
            ):
                return int(ticket_number)
        return None
