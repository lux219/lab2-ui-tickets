from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True, order=True)
class Student:
    last_name: str
    first_name: str

    @property
    def full_name(self) -> str:
        return f"{self.last_name} {self.first_name}"


class StudentStore:
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"Файл со студентами не найден: {self.path}")

    def group_names(self) -> list[str]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def students_for_group(self, group_name: str) -> list[Student]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if group_name not in workbook.sheetnames:
                return []

            sheet = workbook[group_name]
            students: list[Student] = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                last_name = _clean_cell(row[0] if len(row) > 0 else None)
                first_name = _clean_cell(row[1] if len(row) > 1 else None)
                if last_name and first_name:
                    students.append(Student(last_name=last_name, first_name=first_name))
            return students
        finally:
            workbook.close()


def _clean_cell(value: object) -> str:
    return "" if value is None else str(value).strip()
